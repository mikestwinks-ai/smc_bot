from __future__ import annotations

import json
import sqlite3
from types import SimpleNamespace
import pandas as pd
import pytest
from openpyxl import load_workbook

from smc.execution.fills import exec_key
from smc.core.exceptions import MT5DataError
from smc.journal.export import export_excel
from smc.journal.monitoring import MonitoringService
from smc.journal.reconciliation import MT5Reconciler
from smc.journal.store import JournalStore
from tests.harness import add_plan, make_env


def test_migrations_preserve_existing_events_and_use_utc(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    db = sqlite3.connect(path)
    db.execute("""CREATE TABLE journal_events (
        event_id TEXT PRIMARY KEY, occurred_at TEXT NOT NULL, event_type TEXT NOT NULL,
        correlation_id TEXT, idempotency_key TEXT UNIQUE, plan_id TEXT, order_ticket INTEGER,
        deal_ticket INTEGER, position_id INTEGER, symbol TEXT, direction TEXT, payload_json TEXT NOT NULL
    )""")
    db.execute("PRAGMA user_version=1")
    db.execute("INSERT INTO journal_events VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
               ("old", "2026-01-01T00:00:00+00:00", "system_start", None, None, None, None, None,
                None, None, None, "{}"))
    db.commit()
    db.close()

    journal = JournalStore(path)
    assert journal._db.execute("PRAGMA user_version").fetchone()[0] == 3
    assert journal.recent_events()[0]["event_id"] == "old"
    journal.record_event("test", occurred_at=pd.Timestamp("2026-01-02"), idempotency_key="same-key")
    journal.record_event("test-duplicate", occurred_at=pd.Timestamp("2026-01-02"), idempotency_key="same-key")
    rows = journal.events()
    assert len([event for event in rows if event["idempotency_key"] == "same-key"]) == 1
    assert rows[-1]["occurred_at"].endswith("+00:00")
    journal.close()


def test_reconciliation_is_idempotent_and_aggregates_partial_costs_by_position(tmp_path):
    env = make_env(tmp_path, now="2026-01-07 10:00")
    plan = add_plan(env)
    position_id = env.broker.add_closed_trade(
        net=5.0, minutes_ago=20, magic=env.cfg.execution.magic_number, partial_split=True,
    )
    key = exec_key(plan.plan_id, plan.created_at)
    for deal in env.broker.deals[-3:]:
        deal.comment = key
        deal.swap = -0.2
        deal.fee = -0.1
    env.broker.add_closed_trade(net=100.0, minutes_ago=10, magic=99999)
    env.broker.add_symbol("EURUSDm", 1.1, 1.2)
    env.broker.add_closed_trade(net=200.0, minutes_ago=5, magic=env.cfg.execution.magic_number,
                                symbol="EURUSDm")
    env.broker.history_orders_get = lambda *args, **kwargs: (
        SimpleNamespace(ticket=501, magic=env.cfg.execution.magic_number, symbol="XAUUSDm", type=2,
                        state=4, time_setup=int(env.now.timestamp()), time_done=int(env.now.timestamp()),
                        volume_initial=0.1, volume_current=0.0, price_open=115.0, sl=112.0, tp=150.0,
                        comment=key),
        SimpleNamespace(ticket=502, magic=99999, symbol="XAUUSDm", type=2, state=4,
                        time_setup=int(env.now.timestamp()), time_done=int(env.now.timestamp()),
                        volume_initial=0.1, volume_current=0.0, price_open=115.0, sl=112.0, tp=150.0,
                        comment="unrelated"),
    )
    result = env.system.reconcile_history()
    again = env.system.reconcile_history()

    trades = env.system.journal.trades()
    assert result.deals_seen == 3
    assert result.orders_seen == 1
    assert again.deals_seen == 3
    assert len(env.system.journal.broker_orders()) == 1
    assert len(trades) == 1
    trade = trades[0]
    assert trade["position_id"] == position_id
    assert trade["status"] == "closed"
    assert trade["direction"] == "buy"
    assert trade["net"] == pytest.approx(4.1)
    assert trade["costs"] == pytest.approx(-0.9)
    matching = [event for event in env.system.journal.events()
                if event["event_type"] == "deal_reconciled"]
    assert len(matching) == 3
    assert {event["plan_id"] for event in matching} == {plan.plan_id}
    assert {event["correlation_id"] for event in matching} == {key}
    env.system.close()


def test_reconciliation_keeps_open_positions_and_fails_without_partial_commit(tmp_path):
    env = make_env(tmp_path)
    open_position = env.broker.add_open_position(
        "XAUUSDm", True, 0.2, 118.0, 112.0, 150.0, env.cfg.execution.magic_number,
    )
    env.broker._close({"position": open_position.ticket, "volume": 0.1, "price": 119.0}, False)
    other = env.broker.add_open_position("XAUUSDm", False, 0.2, 118.0, 122.0, 100.0, 777)
    result = env.system.reconcile_history()
    assert result.positions_seen == 1
    assert env.system.journal.current_positions(1, env.cfg.execution.magic_number)[0]["ticket"] == open_position.ticket
    trade = next(row for row in env.system.journal.trades() if row["position_id"] == open_position.ticket)
    assert trade["status"] == "open"
    assert trade["volume_closed"] == pytest.approx(0.1)
    assert all(row["position_id"] != other.ticket for row in env.system.journal.trades())

    before_runs = len(env.system.journal.reconciliation_runs())
    env.broker.history_orders_get = lambda *args, **kwargs: (_ for _ in ()).throw(ConnectionError("offline"))
    with pytest.raises(MT5DataError, match="history_orders_get raised"):
        env.system.reconcile_history()
    assert len(env.system.journal.reconciliation_runs()) == before_runs
    assert env.system.journal.current_positions(1, env.cfg.execution.magic_number)[0]["ticket"] == open_position.ticket
    env.system.close()


def test_reversal_history_is_not_falsely_counted_as_a_completed_trade(tmp_path):
    env = make_env(tmp_path)
    position_id = env.broker.add_closed_trade(net=3.0, minutes_ago=10,
                                              magic=env.cfg.execution.magic_number)
    env.broker.deals[-1].entry = 2
    env.broker.deals[-1].volume = 0.2
    result = env.system.reconcile_history()
    trade = next(row for row in env.system.journal.trades() if row["position_id"] == position_id)
    assert result.trades_closed == 0
    assert trade["status"] == "unknown"
    env.system.close()


def test_excel_export_includes_required_sheets_and_never_overwrites(tmp_path):
    journal = JournalStore(tmp_path / "journal.sqlite3")
    journal.record_event("news_lock_on", occurred_at=pd.Timestamp("2026-01-01", tz="UTC"),
                         idempotency_key="news-lock", symbol="XAUUSD")
    output = tmp_path / "exports" / "phase6.xlsx"
    assert export_excel(journal, output) == output
    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == [
        "Trade Summary", "Setup & Plan", "Orders & Fills", "Risk & Rejections", "News Events", "System Health",
    ]
    assert workbook["Trade Summary"].max_row == 1
    assert workbook["News Events"].max_row == 2
    workbook.close()
    with pytest.raises(FileExistsError):
        export_excel(journal, output)
    journal.close()


def test_journal_write_failure_does_not_bypass_execution_guards(tmp_path, monkeypatch):
    env = make_env(tmp_path)

    def fail_write(*args, **kwargs):
        raise sqlite3.OperationalError("simulated database interruption")

    monkeypatch.setattr(env.system.journal, "record_event", fail_write)
    result = env.system.engine.process_plan(add_plan(env))
    assert result.outcome.value == "dry_run"
    assert env.broker.sent == []
    assert env.broker.checked
    env.system.close()


def test_monitoring_is_read_only_and_does_not_persist_daily_baseline(tmp_path):
    env = make_env(tmp_path)
    service = MonitoringService(env.cfg, env.system.gateway, env.system.exposure, env.system.risk,
                                env.system.kill, env.mgr.store, env.system.journal)
    day = env.system.clock.trading_day_key(env.now)
    before_sent = list(env.broker.sent)
    status = service.snapshot(env.now)
    assert status["mt5"]["connected"] is True
    assert status["account"]["login"] == 1
    assert status["execution"]["live_enabled"] is False
    assert status["risk"]["daily_trade_count"] == 0
    assert env.system.state.get_baseline(day) is None
    assert env.broker.sent == before_sent
    env.system.close()
