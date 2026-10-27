"""Deterministic, non-overwriting Excel exports from persisted journal records."""

from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from smc.journal.store import JournalStore


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, sort_keys=True, default=str)
    return value


def _write_sheet(workbook: Workbook, name: str, rows: list[dict[str, Any]], columns: tuple[str, ...]) -> None:
    sheet = workbook.create_sheet(name)
    headers = list(columns)
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([_cell(row.get(key)) for key in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_excel(journal: JournalStore, output_path: Path | str) -> Path:
    """Export all stored records; an existing destination is never overwritten."""
    target = Path(output_path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Excel journal export path must end in .xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    reservation = os.open(target, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    os.close(reservation)
    tmp_path: Path | None = None
    try:
        with journal._lock:
            return _build_export(journal, target, tmp_path)
    except BaseException:
        if tmp_path is not None and tmp_path.exists():
            tmp_path.unlink()
        if target.exists() and target.stat().st_size == 0:
            target.unlink()
        raise


def _build_export(journal: JournalStore, target: Path, tmp_path: Path | None) -> Path:
    try:
        trades = journal.trades()
        events = journal.events()
        orders = journal.broker_orders()
        deals = journal.broker_deals()
        current_orders = journal.all_current_orders()

        plans = [e for e in events if e["event_type"].startswith(("setup_", "plan_"))]
        order_events = [e for e in events if e["event_type"].startswith(("order_", "execution_", "deal_", "position_",
                                                                            "management_"))]
        order_rows = []
        for order in orders:
            order_rows.append({"record_type": "order", "event_time": order.get("created_at"),
                               "order_ticket": order.get("order_ticket"), "symbol": order.get("symbol"),
                               "direction": order.get("direction"), "state": order.get("state"),
                               "volume": order.get("volume"), "details": order.get("payload_json")})
        for row in current_orders:
            details = row["snapshot"]
            order_rows.append({"record_type": "pending_order", "event_time": row["updated_at"],
                               "order_ticket": row["ticket"], "symbol": details.get("symbol"),
                               "direction": "buy" if details.get("is_buy") else "sell",
                               "state": "pending", "volume": details.get("volume"),
                               "price": details.get("price_open"), "stop_loss": details.get("sl"),
                               "take_profit": details.get("tp"), "details": details})
        for deal in deals:
            order_rows.append({"record_type": "deal", "event_time": deal.get("occurred_at"),
                               "order_ticket": deal.get("order_ticket"), "deal_ticket": deal.get("deal_ticket"),
                               "position_id": deal.get("position_id"), "symbol": deal.get("symbol"),
                               "direction": deal.get("direction"), "volume": deal.get("volume"),
                               "price": deal.get("price"), "commission": deal.get("commission"),
                               "swap": deal.get("swap"), "fee": deal.get("fee"), "profit": deal.get("profit"),
                               "net": deal.get("net")})
        for event in order_events:
            order_rows.append(_event_row(event))

        risk = [e for e in events if e["event_type"].startswith(("risk_", "trade_rejected", "plan_rejected",
                                                                   "kill_switch"))]
        news = [e for e in events if e["event_type"].startswith("news_")]
        health = [e for e in events if e["event_type"].startswith(("system_", "reconciliation_", "mt5_"))]

        book = Workbook()
        book.properties.creator = "SMC bot"
        book.properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
        book.properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
        book.remove(book.active)
        _write_sheet(book, "Trade Summary", trades, (
            "account_id", "magic", "position_id", "symbol", "direction", "status", "opened_at", "closed_at",
            "volume_open", "volume_closed", "net", "costs", "updated_at",
        ))
        _write_sheet(book, "Setup & Plan", plans, (
            "occurred_at", "event_type", "correlation_id", "plan_id", "symbol", "direction", "payload",
        ))
        _write_sheet(book, "Orders & Fills", order_rows, (
            "record_type", "event_time", "order_ticket", "deal_ticket", "position_id", "symbol", "direction",
            "state", "volume", "price", "commission", "swap", "fee", "profit", "net", "details", "payload",
        ))
        _write_sheet(book, "Risk & Rejections", [_event_row(e) for e in risk], (
            "occurred_at", "event_type", "correlation_id", "plan_id", "symbol", "direction", "payload",
        ))
        _write_sheet(book, "News Events", [_event_row(e) for e in news], (
            "occurred_at", "event_type", "correlation_id", "symbol", "payload",
        ))
        _write_sheet(book, "System Health", [_event_row(e) for e in health], (
            "occurred_at", "event_type", "correlation_id", "payload",
        ))
        handle, temp_name = tempfile.mkstemp(prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent)
        os.close(handle)
        tmp_path = Path(temp_name)
        book.save(tmp_path)
        os.replace(tmp_path, target)
        return target
    except BaseException:
        if tmp_path is not None and tmp_path.exists():
            tmp_path.unlink()
        raise


def _event_row(event: dict[str, Any]) -> dict[str, Any]:
    row = {key: event.get(key) for key in (
        "occurred_at", "event_type", "correlation_id", "plan_id", "order_ticket", "deal_ticket",
        "position_id", "symbol", "direction",
    )}
    row["payload"] = event.get("payload", {})
    return row
